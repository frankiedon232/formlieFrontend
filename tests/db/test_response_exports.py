"""Response exports (B8): CSV, Excel and the PDF report made in the background, kept 7 days, the Exports page, a
one-time 5-minute download link per download, removal, the audit trail and permissions."""

from __future__ import annotations

import io

import openpyxl
import pytest

from app.modules.public import router as public_router
from tests.conftest import client_for
from tests.db.test_audit_trail import jobs, run_jobs  # noqa: F401  (fixture)
from tests.db.test_auth_flows import app, seeded, sign_in  # noqa: F401  (fixtures)
from tests.db.test_public_forms import publish, session, solved
from tests.db.test_responses import build
from tests.db.test_settings import PASSWORD, new_workspace
from tests.secure_client import SecureClient

pytestmark = pytest.mark.db


async def test_response_exports(app, jobs, monkeypatch) -> None:  # noqa: ANN001, F811
    monkeypatch.setattr(public_router, "MIN_FILL_MS", 0)
    sub, email = await new_workspace(app)
    async with client_for(app, host=f"{sub}.formalie.dev") as http:
        owner = SecureClient(http)
        await sign_in(owner, email, PASSWORD)
        form = await publish(owner, "Signups", build)
        key = form["public_key"]
        for name, mail, plan in (("Ana Silva", "ana@example.org", "basic"), ("Ben Okoro", "=cmd|' /C calc'!A0", "pro")):
            visitor = SecureClient(http)
            data = {"full_name": name, "plan": plan, **({"email": mail} if "@" in mail else {"notes": mail})}
            status, reply = await visitor.request("POST", f"/public/forms/{key}/submit", body={
                "data": data, "challenge": await solved(visitor, key)}, headers=session())
            assert status == 201, reply

        _, reply = await owner.request("POST", f"/forms/{form['id']}/responses/export", body={
            "format": "csv", "scope": "selected", "ids": ["nope"]})
        assert reply["error"]["details"] == [{"field": "scope", "message": "Nothing to export."}]

        made = {}
        for fmt in ("csv", "xlsx", "pdf"):
            status, reply = await owner.request("POST", f"/forms/{form['id']}/responses/export", body={
                "format": fmt, "scope": "filtered", "query": {"filter[a_plan]": "pro"} if fmt == "csv" else {}})
            assert status == 202, reply
            made[fmt] = reply["data"]
            assert reply["data"]["status"] == "queued" and reply["data"]["file_name"].endswith(f".{fmt}")
        assert made["csv"]["rows"] == 1 and made["xlsx"]["rows"] == 2 and made["csv"]["columns"] == 5

        _, reply = await owner.request("POST", f"/responses/exports/{made['csv']['id']}/link", body={})
        assert reply["error"]["details"] == [{"field": "status", "message": "The file is still being made."}]
        await run_jobs(jobs)

        _, reply = await owner.request("GET", "/responses/exports")
        assert {e["status"] for e in reply["data"]} == {"ready"} and reply["meta"]["total"] == 3
        _, reply = await owner.request("GET", "/responses/exports", query={"filter[format]": "pdf"})
        assert [e["format"] for e in reply["data"]] == ["pdf"]

        async def download(export: dict) -> bytes:
            _, reply = await owner.request("POST", f"/responses/exports/{export['id']}/link", body={})
            url = reply["data"]["url"]
            first = await http.get(url)
            assert first.status_code == 200, first.text
            assert (await http.get(url)).status_code == 410  # one use
            return first.content

        csv = (await download(made["csv"])).decode("utf-8-sig")
        assert csv.splitlines()[0].startswith("Number,Submitted (UTC),Status,Tags,Came in,Language")
        assert "'=cmd|' /C calc'!A0" in csv and "Ana Silva" not in csv  # formula defused; filtered to "pro"
        book = openpyxl.load_workbook(io.BytesIO(await download(made["xlsx"])))
        sheet = book.active
        assert sheet.title == "Signups" and sheet.cell(1, 1).value == "Number" and sheet.max_row == 3
        pdf = await download(made["pdf"])
        assert pdf.startswith(b"%PDF-1.4") and b"Signups" in pdf and b"Page 1 of 1" in pdf

        # Removal; the audit trail
        _, reply = await owner.request("DELETE", f"/responses/exports/{made['csv']['id']}")
        assert reply["data"] == {"id": made["csv"]["id"]}
        _, reply = await owner.request("GET", f"/responses/exports/{made['csv']['id']}")
        assert reply["error"]["code"] == "FRM-GEN-1004"
        _, trail = await owner.request("GET", "/audit-logs", query={"page_size": "100"})
        actions = {row["action"] for row in trail["data"]}
        assert {"responses.exported", "responses.export_downloaded", "responses.export_deleted"} <= actions

    other, other_email = await new_workspace(app)
    async with client_for(app, host=f"{other}.formalie.dev") as http:
        stranger = SecureClient(http)
        await sign_in(stranger, other_email, PASSWORD)
        _, reply = await stranger.request("GET", f"/responses/exports/{made['pdf']['id']}")
        assert reply["error"]["code"] == "FRM-GEN-1004"
